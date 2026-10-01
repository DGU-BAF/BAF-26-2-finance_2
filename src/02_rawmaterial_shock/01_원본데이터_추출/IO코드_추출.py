# --- 경로 설정: 레포 어디에서 실행해도 data/ 폴더를 찾아 루트로 지정 (절대경로 사용 금지) ---
from pathlib import Path as _P
ROOT = next(p for p in _P(__file__).resolve().parents if (p / "data" / "raw").exists())
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"
SHOCK_INTERIM = INTERIM / "원자재_충격지표"
PROCESSED = ROOT / "data" / "processed"
SHOCK_INTERIM.mkdir(parents=True, exist_ok=True)
PROCESSED.mkdir(parents=True, exist_ok=True)
# ---------------------------------------------------------------------------------
import openpyxl
from pathlib import Path
import csv

BASE = ROOT
path = RAW / "원자재" / "투입산출표_생산자가격_기본부문_2019년연장표.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb["국산투입계수(Ad)"]

# row-side codes/names: column A/B, rows 7..388
out_rows = SHOCK_INTERIM / "IO코드_행목록.csv"
with open(out_rows, "w", newline="", encoding="utf-8-sig") as f:
    writer = csv.writer(f)
    writer.writerow(["code", "name"])
    for row in ws.iter_rows(min_row=7, max_row=388, max_col=2, values_only=True):
        code, name = row
        if code is None:
            continue
        writer.writerow([code, name])

# column-side codes/names: row 5 (codes), row 6 (names), columns C..386
codes_row = next(ws.iter_rows(min_row=5, max_row=5, min_col=3, max_col=386, values_only=True))
names_row = next(ws.iter_rows(min_row=6, max_row=6, min_col=3, max_col=386, values_only=True))
out_cols = SHOCK_INTERIM / "IO코드_열목록.csv"
with open(out_cols, "w", newline="", encoding="utf-8-sig") as f:
    writer = csv.writer(f)
    writer.writerow(["code", "name"])
    for code, name in zip(codes_row, names_row):
        if code is None:
            continue
        writer.writerow([code, name])

print("row codes written:", out_rows)
print("col codes written:", out_cols)

# quick counts
with open(out_rows, encoding="utf-8-sig") as f:
    print("row count:", sum(1 for _ in f) - 1)
with open(out_cols, encoding="utf-8-sig") as f:
    print("col count:", sum(1 for _ in f) - 1)
