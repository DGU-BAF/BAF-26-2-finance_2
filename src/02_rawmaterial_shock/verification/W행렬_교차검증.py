# --- 경로 설정: 레포 어디에서 실행해도 data/ 폴더를 찾아 루트로 지정 (절대경로 사용 금지) ---
from pathlib import Path as _P
ROOT = next(p for p in _P(__file__).resolve().parents if (p / "data" / "raw").exists())
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"
SHOCK_INTERIM = INTERIM / "원자재_충격지표"
PROCESSED = ROOT / "data" / "processed"
SHOCK_INTERIM.mkdir(parents=True, exist_ok=True)
PROCESSED.mkdir(parents=True, exist_ok=True)
# ---------------------------------------------------------------------------------
import openpyxl
from pathlib import Path
import csv

w_path = INTERIM / "가중치행렬_W_2019_검증본.xlsx"
our_xwalk_path = SHOCK_INTERIM / "최종_크로스워크표.csv"

wb = openpyxl.load_workbook(w_path, read_only=True, data_only=True)

# Industries 시트 전체 - 40개 업종 다 있는지, W행합 전부 1인지
ws_ind = wb["Industries"]
rows = list(ws_ind.iter_rows(min_row=5, max_row=ws_ind.max_row, max_col=5, values_only=True))
rows = [r for r in rows if r[0] is not None]
print("Industries 시트 업종 수:", len(rows))
bad = [r for r in rows if r[4] is not None and abs(r[4]-1) > 1e-6]
print("W 행합이 1이 아닌 업종 수:", len(bad))
for r in bad[:10]:
    print(r)

# W 시트 크기 확인
ws_w = wb["W"]
print("\nW 시트 max_row:", ws_w.max_row, "max_col:", ws_w.max_column)

# Products 시트: 우리 final_crosswalk.csv와 산업코드 일치하는지 대조
our_map = {}
with open(our_xwalk_path, encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        our_map[r["code"]] = r["istans_industry_code"]

ws_p = wb["Products"]
mismatch = []
total = 0
for row in ws_p.iter_rows(min_row=5, max_row=ws_p.max_row, max_col=6, values_only=True):
    code = row[0]
    if code is None:
        continue
    total += 1
    their_ind = row[2]
    our_ind = our_map.get(str(code), "")
    if str(their_ind) != str(our_ind):
        mismatch.append((code, row[1], our_ind, their_ind))

print(f"\n총 {total}개 상품 코드 대조, 불일치 {len(mismatch)}개")
for m in mismatch[:20]:
    print(m)
