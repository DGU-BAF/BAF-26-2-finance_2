# --- 경로 설정: 레포 어디에서 실행해도 data/ 폴더를 찾아 루트로 지정 (절대경로 사용 금지) ---
from pathlib import Path as _P
ROOT = next(p for p in _P(__file__).resolve().parents if (p / "data" / "raw").exists())
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"
SHOCK_INTERIM = INTERIM / "원자재_충격지표"
PROCESSED = ROOT / "data" / "processed"
SHOCK_INTERIM.mkdir(parents=True, exist_ok=True)
PROCESSED.mkdir(parents=True, exist_ok=True)
# ---------------------------------------------------------------------------------
import openpyxl
from pathlib import Path

path = RAW / "원자재" / "투입산출표_생산자가격_기본부문_2019년연장표.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

ws = wb["국산투입계수(Ad)"]
print("max_row:", ws.max_row, "max_col:", ws.max_column)

rows = list(ws.iter_rows(min_row=1, max_row=10, max_col=4, values_only=True))
for r in rows:
    print(r)

print("---last few rows/cols (totals) ---")
rows2 = list(ws.iter_rows(min_row=383, max_row=388, max_col=4, values_only=True))
for r in rows2:
    print(r)
