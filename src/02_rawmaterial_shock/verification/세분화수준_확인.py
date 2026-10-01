# --- 경로 설정: 레포 어디에서 실행해도 data/ 폴더를 찾아 루트로 지정 (절대경로 사용 금지) ---
from pathlib import Path as _P
ROOT = next(p for p in _P(__file__).resolve().parents if (p / "data" / "raw").exists())
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"
SHOCK_INTERIM = INTERIM / "원자재_충격지표"
PROCESSED = ROOT / "data" / "processed"
SHOCK_INTERIM.mkdir(parents=True, exist_ok=True)
PROCESSED.mkdir(parents=True, exist_ok=True)
# ---------------------------------------------------------------------------------
import openpyxl
from pathlib import Path

path = RAW / "원자재" / "ISTANS_60대산업_산업연관_연계표.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb["2015년기준"]

targets = {"3102", "4011", "3308", "3312", "2000", "0111"}
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=10, values_only=True):
    iok4, iok4n, sic, sicn = row[6], row[7], row[8], row[9]
    if iok4 in targets:
        print(row)

print("\n--- unique industry codes count & sample ---")
codes = set()
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=10, values_only=True):
    codes.add((row[8], row[9]))
print("unique 산업코드 count:", len(codes))
for c in sorted(codes, key=lambda x: (len(x[0]), x[0]))[:70]:
    print(c)
