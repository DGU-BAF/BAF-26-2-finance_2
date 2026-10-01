# --- 경로 설정: 레포 어디에서 실행해도 data/ 폴더를 찾아 루트로 지정 (절대경로 사용 금지) ---
from pathlib import Path as _P
ROOT = next(p for p in _P(__file__).resolve().parents if (p / "data" / "raw").exists())
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"
SHOCK_INTERIM = INTERIM / "원자재_충격지표"
PROCESSED = ROOT / "data" / "processed"
SHOCK_INTERIM.mkdir(parents=True, exist_ok=True)
PROCESSED.mkdir(parents=True, exist_ok=True)
# ---------------------------------------------------------------------------------
import openpyxl
from pathlib import Path

path = INTERIM / "가중치행렬_W_2019_검증본.xlsx"
print("파일 존재:", path.exists())
if not path.exists():
    raise SystemExit

wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("시트 목록:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== {sheet_name} === max_row={ws.max_row} max_col={ws.max_column}")
    rows = list(ws.iter_rows(min_row=1, max_row=5, max_col=6, values_only=True))
    for r in rows:
        print(r)
