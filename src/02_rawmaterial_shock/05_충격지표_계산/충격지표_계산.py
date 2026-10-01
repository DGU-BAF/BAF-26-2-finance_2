# --- 경로 설정: 레포 어디에서 실행해도 data/ 폴더를 찾아 루트로 지정 (절대경로 사용 금지) ---
from pathlib import Path as _P
ROOT = next(p for p in _P(__file__).resolve().parents if (p / "data" / "raw").exists())
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"
SHOCK_INTERIM = INTERIM / "원자재_충격지표"
PROCESSED = ROOT / "data" / "processed"
SHOCK_INTERIM.mkdir(parents=True, exist_ok=True)
PROCESSED.mkdir(parents=True, exist_ok=True)
# ---------------------------------------------------------------------------------
# -*- coding: utf-8 -*-
"""
원자재 가격충격 지표(RawShock) 계산 — 최종 파이프라인
Shock_40,t = W @ (M' @ Δp_381,t)

- M  = '수입유발계수' 시트 (직접+간접 수입노출, 이미 BOK가 계산해서 제공, 381x381)
- Δp = 확정된 8개 카테고리 IO코드에 해당하는 ECOS 원자재 가격(원화기준)의
       100*ln(P_t/P_t-1). 매칭 안 된 코드는 0(=이번 분석에서 충격을 안 넣는다는 뜻,
       실제 가격이 안 변했다는 뜻이 아님).
- W  = ISTANS 40개 업종 x 381개 상품 산출액가중평균 행렬(팀원 제공, 교차검증 완료)

기본안: 섬유원료=1111만 사용
확장검증: 섬유원료=1111+1119(1119는 1111 가격을 대리지수로 사용)
"""
import csv
import numpy as np
import openpyxl
from pathlib import Path

BASE = ROOT
ECOS_CSV = RAW / "원자재" / "수입물가지수(기본분류)_26233649.csv"
IO_XLSX = RAW / "원자재" / "투입산출표_생산자가격_기본부문_2019년연장표.xlsx"
IO_CODES_CSV = SHOCK_INTERIM / "IO코드_행목록.csv"
W_CSV = INTERIM / "가중치행렬_W_2019_계산용.csv"
OUT_CSV = PROCESSED / "원자재_가격충격_패널.csv"

# ------------------------------------------------------------
# 1) IO 코드 순서 (381개, 9590 합계행 제외) - 이 순서가 M, W 둘 다의 기준 순서
# ------------------------------------------------------------
io_codes = []
with open(IO_CODES_CSV, encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        if r["code"] != "9590":
            io_codes.append(r["code"])
n = len(io_codes)
code_to_idx = {c: i for i, c in enumerate(io_codes)}
print("IO 코드 수:", n)

# ------------------------------------------------------------
# 2) 확정된 8개 카테고리 -> IO코드 -> ECOS 계정코드(이름) 매핑
#    (final_crosswalk.csv / ecos_io_match_draft.csv에서 이미 확인된 정확매칭 이름 그대로 사용)
# ------------------------------------------------------------
# {IO코드: ECOS 계정코드 이름}
IO_TO_ECOS_NAME = {
    "0621": "원유",
    "2711": "선철", "2712": "합금철", "2713": "조강",
    "2721": "철근및봉강", "2722": "형강", "2724": "중후판(두께3mm이상)",
    "2725": "열연강판", "2727": "철강관", "2730": "냉간압연강재", "2791": "표면처리강재",
    # 2723(선재및궤조), 2726(강선), 2799(기타철강1차제품)는 ECOS에 대응 항목 없음 - 매칭 없음
    "2811": "동제련,정련및합금제품", "2812": "알루미늄제련,정련및합금제품",
    "2813": "연및아연제련,정련및합금제품", "2814": "금은괴",
    "2819": "기타비철금속제련,정련및합금제품",
    "2821": "동1차제품", "2822": "알루미늄1차제품", "2829": "기타비철금속1차제품",
    "0302": "원목", "1311": "제재목", "1312": "합판", "1313": "강화및재생목재",
    "1410": "펄프",
    "0729": "기타비금속광물",
    "1111": "천연및화학섬유사",
    "0112": "맥류및잡곡", "0113": "콩류", "0193": "잎담배", "0194": "천연고무",
    "0196": "기타식용작물", "0199": "기타비식용작물",
}
# 1119는 별도 처리(1111 대리)
EXTENDED_ONLY = {"1119": "천연및화학섬유사"}  # 대리지수: 1111과 같은 ECOS 이름 사용

for code in list(IO_TO_ECOS_NAME) + list(EXTENDED_ONLY):
    assert code in code_to_idx, f"IO코드 {code}가 381개 목록에 없음"

# ------------------------------------------------------------
# 3) ECOS CSV에서 연도별 원화기준 지수값 로드 (계정코드 이름 -> {year:value})
# ------------------------------------------------------------
with open(ECOS_CSV, encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    header = next(reader)
    years = [int(y) for y in header[6:]]
    ecos_rows = [r for r in reader if r[2] == "원화기준"]

ecos_series = {}
for r in ecos_rows:
    name = r[1].strip()
    vals = {}
    for y, v in zip(years, r[6:]):
        try:
            vals[y] = float(v)
        except ValueError:
            vals[y] = None
    # 같은 이름이 여러 번 나오면(트리 중복) 첫 번째만 사용
    if name not in ecos_series:
        ecos_series[name] = vals

for name in set(IO_TO_ECOS_NAME.values()) | set(EXTENDED_ONLY.values()):
    assert name in ecos_series, f"ECOS 항목 '{name}' 못 찾음"

# ------------------------------------------------------------
# 4) 연도별 Δp(=100*ln변화율) 계산, 두 버전(기본안/확장검증)
# ------------------------------------------------------------
def log_changes(name):
    s = ecos_series[name]
    out = {}
    for y in years:
        if y-1 in s and y in s and s[y-1] not in (None, 0) and s[y] is not None:
            out[y] = 100 * np.log(s[y] / s[y-1])
    return out

dp_by_code_baseline = {code: log_changes(name) for code, name in IO_TO_ECOS_NAME.items()}
dp_by_code_extended = dict(dp_by_code_baseline)
for code, name in EXTENDED_ONLY.items():
    dp_by_code_extended[code] = log_changes(name)  # 1111과 동일 계산 -> 사실상 대리

analysis_years = [y for y in years if y >= 2015]

def build_dp_vector(dp_by_code, year):
    v = np.zeros(n)
    for code, series in dp_by_code.items():
        if year in series:
            v[code_to_idx[code]] = series[year]
    return v

# ------------------------------------------------------------
# 5) M(수입유발계수) 로드, 381x381, io_codes 순서로 정렬
# ------------------------------------------------------------
wb = openpyxl.load_workbook(IO_XLSX, read_only=True, data_only=True)
ws = wb["수입유발계수"]

# 행 순서(A열, 7~387행)와 io_codes 순서가 같은지 확인 후 그대로 읽기
row_codes = [str(r[0]) for r in ws.iter_rows(min_row=7, max_row=387, max_col=1, values_only=True)]
assert row_codes == io_codes, "행 순서가 io_codes와 다름 - 재확인 필요"

col_codes = [str(v) for v in next(ws.iter_rows(min_row=5, max_row=5, min_col=3, max_col=383, values_only=True))]
assert col_codes == io_codes, "열 순서가 io_codes와 다름 - 재확인 필요"

M = np.zeros((n, n))
for i, row in enumerate(ws.iter_rows(min_row=7, max_row=387, min_col=3, max_col=383, values_only=True)):
    M[i, :] = [x if x is not None else 0.0 for x in row]

print("M 행렬 shape:", M.shape, "결측 0 대체 완료")

# ------------------------------------------------------------
# 6) W(40x381) 로드, 열 순서 io_codes와 동일한지 확인
# ------------------------------------------------------------
with open(W_CSV, encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    w_header = next(reader)
    w_rows = list(reader)

assert w_header[1:] == io_codes, "W 열 순서가 io_codes와 다름"
industry_codes = [r[0] for r in w_rows]
W = np.array([[float(x) for x in r[1:]] for r in w_rows])
print("W 행렬 shape:", W.shape, "업종 수:", len(industry_codes))

# ------------------------------------------------------------
# 7) 연도별 계산: Shock_381 = M^T @ Δp ; Shock_40 = W @ Shock_381
# ------------------------------------------------------------
out_rows = []
for year in analysis_years:
    dp_base = build_dp_vector(dp_by_code_baseline, year)
    dp_ext = build_dp_vector(dp_by_code_extended, year)

    shock381_base = M.T @ dp_base
    shock381_ext = M.T @ dp_ext

    shock40_base = W @ shock381_base
    shock40_ext = W @ shock381_ext

    for g_idx, ind_code in enumerate(industry_codes):
        out_rows.append({
            "year": year,
            "industry_code": ind_code,
            "shock_baseline_1111only": shock40_base[g_idx],
            "shock_extended_1111_1119": shock40_ext[g_idx],
            "diff_1119_effect": shock40_ext[g_idx] - shock40_base[g_idx],
        })

with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
    w_ = csv.DictWriter(f, fieldnames=list(out_rows[0].keys()))
    w_.writeheader()
    w_.writerows(out_rows)

print(f"\n저장 완료: {OUT_CSV} (연도 {analysis_years[0]}~{analysis_years[-1]} x 업종 {len(industry_codes)}개)")

# 2022년 기준 상위 5개 업종 미리보기
print("\n--- 2022년 기준 기본안 충격 상위 5개 업종 ---")
rows_2022 = [r for r in out_rows if r["year"] == 2022]
rows_2022_sorted = sorted(rows_2022, key=lambda r: -abs(r["shock_baseline_1111only"]))
for r in rows_2022_sorted[:5]:
    print(r)
