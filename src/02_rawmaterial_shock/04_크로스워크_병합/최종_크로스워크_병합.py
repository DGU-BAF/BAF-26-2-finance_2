# --- 경로 설정: 레포 어디에서 실행해도 data/ 폴더를 찾아 루트로 지정 (절대경로 사용 금지) ---
from pathlib import Path as _P
ROOT = next(p for p in _P(__file__).resolve().parents if (p / "data" / "raw").exists())
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"
SHOCK_INTERIM = INTERIM / "원자재_충격지표"
PROCESSED = ROOT / "data" / "processed"
SHOCK_INTERIM.mkdir(parents=True, exist_ok=True)
PROCESSED.mkdir(parents=True, exist_ok=True)
# ---------------------------------------------------------------------------------
# -*- coding: utf-8 -*-
"""
1번(원자재 카테고리 태깅) + 2번(ISTANS 40개 업종 매핑, ISTANS 공식 연계표 사용)을
하나의 표로 합친다. 2015년기준 시트 사용(2019년 연장표와 분류체계가 맞을 것으로 추정 -
100% 확정 아님, 팀 확인 필요).
"""
import csv
import openpyxl
from pathlib import Path

BASE = ROOT
xwalk_path = RAW / "원자재" / "ISTANS_60대산업_산업연관_연계표.xlsx"
tagging_path = SHOCK_INTERIM / "원자재_카테고리_태깅표.csv"
out_path = SHOCK_INTERIM / "최종_크로스워크표.csv"

# 1) ISTANS 공식 연계표 로드 (IOK4코드 -> 산업코드/산업명)
wb = openpyxl.load_workbook(xwalk_path, read_only=True, data_only=True)
ws = wb["2015년기준"]
istans_map = {}
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=10, values_only=True):
    iok4 = row[6]
    if iok4 is None:
        continue
    istans_map[str(iok4)] = {"산업코드": row[8], "산업명": row[9], "IOK1N": row[1], "IOK2N": row[3]}

# 2) 1번 태깅 결과 로드
tagging = {}
with open(tagging_path, encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        tagging[r["code"]] = r

# 3) 병합 - tagging의 순서(원본 381개 순서) 기준으로 순회
rows_out = []
unmatched = []
for code, t in tagging.items():
    im = istans_map.get(code)
    if im is None:
        unmatched.append(code)
        istans_code, istans_name = "", ""
    else:
        istans_code, istans_name = im["산업코드"], im["산업명"]
    rows_out.append({
        "code": code,
        "name": t["name"],
        "raw_material_category": t["category"],
        "raw_material_confidence": t["confidence"],
        "raw_material_note": t["note"],
        "istans_industry_code": istans_code,
        "istans_industry_name": istans_name,
    })

with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
    writer = csv.DictWriter(f, fieldnames=list(rows_out[0].keys()))
    writer.writeheader()
    writer.writerows(rows_out)

print("총 행:", len(rows_out))
print("ISTANS 매핑 안 된 코드 수:", len(unmatched))
if unmatched:
    print("미매핑 코드:", unmatched[:20])

# 40개 제조업(I31~I34)에 속하면서 원자재 태깅도 있는 행 = 우리가 최종적으로 쓸 후보
both = [r for r in rows_out if r["raw_material_category"] not in ("해당없음",) and str(r["istans_industry_code"]).startswith("I3")]
print("\n원자재 태깅 있고 + 40개 제조업으로 매핑된 코드 수:", len(both))
for r in both:
    print(r["code"], r["name"], "->", r["raw_material_category"], "/", r["istans_industry_code"], r["istans_industry_name"])
